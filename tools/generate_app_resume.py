"""Builds a .docx resume from one row of the applications table -- the exact
headline, summary, skills, job history and education the bot generated for
that application.

Usage:  python tools/generate_app_resume.py <application id>
"""
import ast
import re
import sqlite3
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DB_PATH = ROOT / "IndHelperDB.db"
OUT_DIR = ROOT / "Resumes Out"

# Not stored per application; same contact line the cover letters use.
CONTACT = ("(770) 383-5362  |  tommyorok43@gmail.com  |  Powder Springs, GA  |  "
           "linkedin.com/in/tommy-orok-8273ab120")

MONTHS = {m: i for i, m in enumerate(
    ["january", "february", "march", "april", "may", "june", "july", "august",
     "september", "october", "november", "december"], 1)}


def load_application(app_id):
    conn = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    row = conn.execute("SELECT * FROM applications WHERE id = ?", (app_id,)).fetchone()
    conn.close()
    if row is None:
        sys.exit(f"No application with id {app_id}")
    return dict(row)


def parse_list(text):
    # Stored as Python reprs (single quotes), not JSON.
    return ast.literal_eval(text) if text else []


def date_key(text):
    """Sort key for 'October 2017' / '2019' style dates."""
    parts = (text or "").lower().split()
    year = int(next((p for p in parts if p.isdigit()), 0))
    month = next((MONTHS[p] for p in parts if p in MONTHS), 0)
    return year, month


def bullets(desc):
    lines = [re.sub(r"^[-•*]\s*", "", line).strip() for line in (desc or "").split("\n")]
    return [line for line in lines if line]


def rule_below(paragraph):
    """Thin line under a section heading."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, val in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "808080")):
        bottom.set(qn(f"w:{key}"), val)
    borders.append(bottom)
    p_pr.append(borders)


def build(app, out_path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(2)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.5)
        section.left_margin = section.right_margin = Inches(0.7)

    def heading(text):
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        rule_below(p)

    def bullet(text, keep_with_next=False):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = keep_with_next

    # --- Header ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(app["fullName"])
    run.bold = True
    run.font.size = Pt(20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(app["headline"])
    run.italic = True
    run.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(CONTACT).font.size = Pt(9.5)

    # --- Summary ---
    heading("Summary")
    p = doc.add_paragraph(app["resumeSummary"].strip())
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # --- Skills ---
    heading("Skills")
    # Two borderless bullet columns: a single "a • b • c" line wraps with a
    # stray bullet starting the next line.
    skills = parse_list(app["skills"])
    half = (len(skills) + 1) // 2
    table = doc.add_table(rows=half, cols=2)
    for i, skill in enumerate(skills):
        cell = table.cell(i % half, i // half)
        cell.paragraphs[0].add_run(f"•  {skill}")
        cell.paragraphs[0].paragraph_format.space_after = Pt(1)

    # --- Experience, most recent first ---
    heading("Experience")
    jobs = sorted(parse_list(app["jobHist"]), key=lambda j: date_key(j.get("fromDate")), reverse=True)
    for job in jobs:
        end = "Present" if job.get("current") == "Yes" else job.get("toDate", "")
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True  # never strand a job header at a page bottom
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.1), alignment=2)  # right-aligned
        title = p.add_run(job["title"])
        title.bold = True
        p.add_run(f"  |  {job['comp']}")
        p.add_run(f"\t{job.get('fromDate', '')} – {end}").italic = True

        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        loc = p.add_run(job.get("cityState", ""))
        loc.italic = True
        loc.font.size = Pt(9.5)
        loc.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

        lines = bullets(job.get("desc"))
        for i, line in enumerate(lines):
            bullet(line, keep_with_next=i < len(lines) - 1)

    # --- Education ---
    heading("Education")
    for edu in parse_list(app["eduHist"]):
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.1), alignment=2)
        run = p.add_run(f"{edu['educationLevel']}, {edu['fieldOfStudy']}")
        run.bold = True
        p.add_run(f"  |  {edu['school']}, {edu.get('cityState', '')}")
        p.add_run(f"\t{edu.get('frmDate', '')} – {edu.get('toDate', '')}").italic = True

    doc.save(out_path)


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    app = load_application(int(sys.argv[1]))
    safe_company = re.sub(r"[^A-Za-z0-9]+", "_", app["companyName"]).strip("_")
    name = app["fullName"].replace(" ", "_")
    OUT_DIR.mkdir(exist_ok=True)
    out_path = OUT_DIR / f"{name}_Resume_{safe_company}.docx"
    build(app, out_path)
    print(out_path)


if __name__ == "__main__":
    main()
